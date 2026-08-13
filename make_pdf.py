from fpdf import FPDF
import docx

# Read the docx
doc = docx.Document(r'labwork\AI_Integration_VNTU.docx')

pdf = FPDF()
pdf.set_auto_page_break(auto=True, margin=15)
pdf.add_page()
pdf.add_font('Arial', '', r'C:\Windows\Fonts\arial.ttf')
pdf.add_font('Arial', 'B', r'C:\Windows\Fonts\arialbd.ttf')
pdf.add_font('Arial', 'I', r'C:\Windows\Fonts\ariali.ttf')
pdf.add_font('Arial', 'BI', r'C:\Windows\Fonts\arialbi.ttf')

pdf.set_font('Arial', size=14)

for p in doc.paragraphs:
    alignment = p.alignment
    text = p.text
    
    if not text.strip():
        if alignment == 1:  # Center
            pdf.ln(10)
        continue
    
    if alignment == 1:  # CENTER
        pdf.ln(5)
        pdf.set_x((210 - pdf.get_string_width(text)) / 2)
        if p.runs and p.runs[0].bold:
            pdf.set_font('Arial', 'B', 14)
        else:
            pdf.set_font('Arial', '', 14)
        pdf.cell(pdf.get_string_width(text), 8, text, align='C')
        pdf.ln(8)
    elif alignment == 3:  # JUSTIFY
        pdf.set_font('Arial', '', 14)
        pdf.multi_cell(0, 8, text, align='J')
    else:
        pdf.set_font('Arial', '', 14)
        pdf.cell(pdf.get_string_width(text) + 5, 8, text)

# Add references
pdf.ln(5)
pdf.set_font('Arial', 'B', 14)
pdf.cell(0, 8, 'Література', align='C')
pdf.ln(5)
pdf.set_font('Arial', '', 14)

references = [
    'Человеко-машинне взаємодія в освіті: Навчальний посібник / Загальнодоступні джерела. – Київ, 2024.',
    'Інтеграція штучного інтелекту в навчальні системи / А. І. Петренко. – Вінниця: ВНЦКІ, 2023.',
    'Методи штучного інтелекту в педагогічній практиці / С. В. Коваль. – Львів, 2022.',
    'Основи штучного інтелекту для викладачів / О. О. Шевченко. – Київський національний університет, 2023.',
    'AI в освіті: можливості та обмеження / Д. М. Іванов. – Харківський національний університет, 2024.'
]

for i, ref in enumerate(references, 1):
    pdf.set_x(15)
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(10, 6, str(i))
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 6, ref)
    pdf.ln()

pdf.output(r'labwork\AI_Integration_VNTU.pdf')
print('PDF created successfully')